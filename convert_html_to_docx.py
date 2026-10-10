#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将 doc-formatter 产出的 HTML 高保真转换为 .docx (A4, portrait, 默认页边距)。

说明：通用库 html2docx 存在 self.r 跨块泄漏的缺陷，会导致标题文本错位、<div> 段落合并，
因此这里使用针对本论文 HTML 结构编写的专用转换器（基于 python-docx + lxml），
并在转换后施加中文学术版式（宋体/黑体、正文两端对齐+首行缩进、各级标题字号、
表格边框、章节分页）。
"""
import os
import sys
from lxml import html as lhtml

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HTML_PATH = r"C:\Users\鄂一铭\WorkBuddy\期中作业2\nlp-digital-transformation-paper\output\aa0664b6-8e10-443d-9faa-16f919183044\stage2\formatted-academic-paper.html"
DOCX_PATH = r"C:\Users\鄂一铭\WorkBuddy\期中作业2\nlp-digital-transformation-paper\paper\基于上市公司年报文本的数字化转型程度测度研究.docx"

TITLE = "基于上市公司年报文本的数字化转型程度测度研究"


def set_eastasia_font(target, ascii_font, ea_font):
    """target: 样式对象或 Run 对象，同时设置西文与中文字体。"""
    target.font.name = ascii_font
    rpr = target._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), ea_font)
    rfonts.set(qn('w:ascii'), ascii_font)
    rfonts.set(qn('w:hAnsi'), ascii_font)


def add_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), '000000')
        borders.append(e)
    tblPr.append(borders)


def add_page_break_before(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:pageBreakBefore'))


class PaperConverter:
    def __init__(self):
        self.doc = Document()
        # 移除 Document() 自带的空段落
        if self.doc.paragraphs and not self.doc.paragraphs[0].text.strip():
            el = self.doc.paragraphs[0]._element
            el.getparent().remove(el)
        self.cur = None          # 当前段落
        self.run = None          # 当前 run
        self.pending_style = None  # 待创建段落的样式（惰性创建，避免空段落）
        self.fmt = set()         # 当前行内格式: b/i/u/code/sub/sup
        self.list_stack = []     # 'number' / 'bullet'
        self.inside_li = 0
        self.tables = []

    # ---- 文本/run ----
    def _get_para(self):
        if self.cur is None:
            self.cur = self.doc.add_paragraph(style=self.pending_style)
            self.run = None
        return self.cur

    def add_text(self, text):
        if not text or not text.strip():
            return
        p = self._get_para()
        r = p.add_run(text)
        self.run = r
        f = r.font
        if 'b' in self.fmt:
            f.bold = True
        if 'i' in self.fmt:
            f.italic = True
        if 'u' in self.fmt:
            f.underline = True
        if 'code' in self.fmt:
            f.name = 'Consolas'
        if 'sub' in self.fmt:
            f.subscript = True
        if 'sup' in self.fmt:
            f.superscript = True

    def new_para(self, style=None):
        self.cur = self.doc.add_paragraph(style=style)
        self.run = None
        return self.cur

    def text_and_children(self, el):
        if el.text:
            self.add_text(el.text)
        for c in el:
            self.walk(c)
            if c.tail:
                self.add_text(c.tail)

    # ---- 主递归 ----
    def walk(self, el):
        tag = el.tag.lower() if isinstance(el.tag, str) else None
        if tag is None:
            return

        # 块级标题
        if tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            lvl = int(tag[1])
            self.cur = None
            self.pending_style = f'Heading {lvl}'
            self.text_and_children(el)
            self.pending_style = None
            self.cur = None
            return

        # 普通段落 / div（不在列表内时产生分段）
        if tag in ('p', 'div') and self.inside_li == 0:
            self.cur = None
            self.pending_style = None
            self.text_and_children(el)
            self.pending_style = None
            self.cur = None
            return

        # 纯容器
        if tag in ('section', 'article', 'main', 'header', 'footer', 'body', 'html', 'nav'):
            self.text_and_children(el)
            return

        if tag == 'br':
            if self.run is not None:
                self.run.add_break()
            else:
                p = self._get_para()
                r = p.add_run()
                r.add_break()
                self.run = r
            return

        if tag == 'hr':
            self.cur = None
            p = self.new_para()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '999999')
            pbdr.append(bottom)
            pPr.append(pbdr)
            self.cur = None
            return

        if tag in ('ul', 'ol'):
            kind = 'bullet' if tag == 'ul' else 'number'
            self.list_stack.append(kind)
            self.text_and_children(el)
            self.list_stack.pop()
            return

        if tag == 'li':
            self.inside_li += 1
            depth = len(self.list_stack)
            kind = self.list_stack[-1] if self.list_stack else 'number'
            base = 'List Bullet' if kind == 'bullet' else 'List Number'
            style_name = base if depth <= 1 else f'{base} {depth}'
            self.cur = None
            self.pending_style = style_name
            self.text_and_children(el)
            self.pending_style = None
            self.cur = None
            self.inside_li -= 1
            return

        if tag == 'table':
            self.cur = None
            self.build_table(el)
            self.cur = None
            return

        if tag in ('thead', 'tbody', 'tr'):
            self.text_and_children(el)
            return

        if tag in ('td', 'th'):
            return  # 由 build_table 处理

        # 行内格式
        if tag in ('b', 'strong'):
            self.fmt.add('b'); self.text_and_children(el); self.fmt.discard('b'); return
        if tag in ('i', 'em'):
            self.fmt.add('i'); self.text_and_children(el); self.fmt.discard('i'); return
        if tag == 'u':
            self.fmt.add('u'); self.text_and_children(el); self.fmt.discard('u'); return
        if tag == 'code':
            self.fmt.add('code'); self.text_and_children(el); self.fmt.discard('code'); return
        if tag == 'sub':
            self.fmt.add('sub'); self.text_and_children(el); self.fmt.discard('sub'); return
        if tag == 'sup':
            self.fmt.add('sup'); self.text_and_children(el); self.fmt.discard('sup'); return
        if tag == 'a':
            self.text_and_children(el); return  # 去掉链接，保留文字
        if tag in ('span', 'font', 'small', 'big', 'label', 'em'):
            self.text_and_children(el); return

        # 未知标签：作为行内容器
        self.text_and_children(el)

    # ---- 表格 ----
    def build_table(self, node):
        rows = []

        def collect(n):
            for c in n:
                if not isinstance(c.tag, str):
                    continue
                t = c.tag.lower()
                if t == 'tr':
                    cells = []
                    for cc in c:
                        if isinstance(cc.tag, str) and cc.tag.lower() in ('td', 'th'):
                            cs = int(cc.get('colspan', 1) or 1)
                            rs = int(cc.get('rowspan', 1) or 1)
                            cells.append((cc, cs, rs))
                    rows.append(cells)
                elif t in ('thead', 'tbody'):
                    collect(c)

        collect(node)
        if not rows:
            return
        ncols = max((sum(cs for _, cs, _ in r) for r in rows), default=1)
        nrows = len(rows)
        table = self.doc.add_table(rows=nrows, cols=ncols)
        self.tables.append(table)
        used = [[False] * ncols for _ in range(nrows)]
        for ri, row in enumerate(rows):
            ci = 0
            for cell_elem, cs, rs in row:
                while ci < ncols and used[ri][ci]:
                    ci += 1
                if ci >= ncols:
                    break
                cell = table.cell(ri, ci)
                cell.text = (cell_elem.text_content() or '').strip()
                c2 = min(ncols - 1, ci + cs - 1)
                r2 = min(nrows - 1, ri + rs - 1)
                if c2 > ci or r2 > ri:
                    try:
                        cell.merge(table.cell(r2, c2))
                    except Exception:
                        pass
                for rr in range(ri, ri + rs):
                    for cc2 in range(ci, ci + cs):
                        if rr < nrows and cc2 < ncols:
                            used[rr][cc2] = True
                ci += cs
        for row in table.rows:
            for cell in row.cells:
                for cpar in cell.paragraphs:
                    cpar.alignment = WD_ALIGN_PARAGRAPH.CENTER


def convert(html: str) -> Document:
    tree = lhtml.fromstring(html)
    head = tree.find('head')
    if head is not None:
        head.getparent().remove(head)
    body = tree.find('body') if tree.find('body') is not None else tree
    conv = PaperConverter()
    conv.walk(body)
    return conv.doc


def main():
    with open(HTML_PATH, "r", encoding="utf-8") as f:
        html = f.read()
    doc = convert(html)

    # ---- 页面：A4 / portrait / 默认(1英寸)页边距 ----
    for section in doc.sections:
        section.page_height = Mm(297)
        section.page_width = Mm(210)
        section.orientation = WD_ORIENT.PORTRAIT
        section.top_margin = Mm(25.4)
        section.bottom_margin = Mm(25.4)
        section.left_margin = Mm(25.4)
        section.right_margin = Mm(25.4)

    # ---- 默认字体：宋体(中文) + Times New Roman(西文)，小四(12pt)，1.5 倍行距 ----
    normal = doc.styles['Normal']
    set_eastasia_font(normal, 'Times New Roman', '宋体')
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5

    h1 = doc.styles['Heading 1']
    set_eastasia_font(h1, 'Arial', '黑体')
    h1.font.size = Pt(18); h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

    h2 = doc.styles['Heading 2']
    set_eastasia_font(h2, 'Arial', '黑体')
    h2.font.size = Pt(15); h2.font.bold = True

    h3 = doc.styles['Heading 3']
    set_eastasia_font(h3, 'Arial', '黑体')
    h3.font.size = Pt(13); h3.font.bold = True

    # ---- 定位封面区间（标题 H1 之后，目录 H2 之前） ----
    paras = doc.paragraphs
    cover_start = None
    toc_idx = None
    for i, p in enumerate(paras):
        if p.style.name == 'Heading 1' and cover_start is None:
            cover_start = i
        if p.style.name == 'Heading 2' and p.text.strip() == '目　录' and toc_idx is None:
            toc_idx = i
            break
    cover_region = range(cover_start + 1, toc_idx) if (cover_start is not None and toc_idx is not None) else range(0)

    # ---- 逐段后处理 ----
    for i, p in enumerate(paras):
        st = p.style.name
        text = p.text.strip()
        if st == 'Heading 1':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif st == 'Heading 2':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if text != '目　录':
                add_page_break_before(p)
        elif st == 'Heading 3':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif st == 'Normal':
            if i in cover_region:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ('：' in text or ':' in text) else WD_ALIGN_PARAGRAPH.CENTER
            elif text == '摘　要':
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif len(text) >= 30:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = Pt(24)
                p.paragraph_format.line_spacing = 1.5

    # ---- 表格：中文字体 + 边框 ----
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for cpar in cell.paragraphs:
                    for run in cpar.runs:
                        set_eastasia_font(run, 'Times New Roman', '宋体')
                        run.font.size = Pt(11)
        add_table_borders(table)

    doc.core_properties.title = TITLE
    os.makedirs(os.path.dirname(DOCX_PATH), exist_ok=True)
    doc.save(DOCX_PATH)
    print("SAVED:", DOCX_PATH)
    print("SIZE:", os.path.getsize(DOCX_PATH), "bytes")
    print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print("ERROR:", repr(e))
        import traceback
        traceback.print_exc()
        sys.exit(1)
